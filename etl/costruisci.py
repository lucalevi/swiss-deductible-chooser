#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Insurek — costruzione dei dati statici.

Questo script gira UNA VOLTA L'ANNO, sul computer di chi mantiene il sito, e
non ha niente a che vedere con il sito servito: legge i file pubblicati
dall'Ufficio federale della sanita' pubblica e ne ricava i JSON che il browser
scarichera'. Dopo di lui insurek.lucalevi.com e' HTML, CSS e JavaScript e
nient'altro: nessun server applicativo, nessun database, nessuna chiamata a
terzi mentre qualcuno usa il sito.

    python3 etl/costruisci.py

Legge da  dati-fonte/  e scrive in  sito/dati/ .

I file sorgente (scaricarli a fine settembre, quando l'UFSP approva i premi
dell'anno dopo):

  Prämien_CH.csv                          opendata.swiss, dataset
  Tarife.csv                              "Krankenversicherungsprämien"
  praemienregionen.xlsx                   priminfo.admin.ch/downloads/
  zugelassene-krankenversicherer-*.xlsx   priminfo.admin.ch/downloads/

Perche' i premi vengono da qui e non da uno scraper di Priminfo: sono la
stessa identica fonte. Il calcolatore di Priminfo e' costruito su questo
dataset, che e' pubblicato con licenza di uso libero. Fare lo scraping del
sito federale significherebbe leggere in HTML fragile dei numeri che qui sono
gia' in CSV, con l'obbligo di reggere ogni loro restyling e un carico inutile
su un servizio pubblico.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Manca openpyxl.  pip install openpyxl")


RADICE = Path(__file__).resolve().parent.parent
FONTE = RADICE / "dati-fonte"
USCITA = RADICE / "sito" / "dati"

# Le sei franchigie per adulti, nell'ordine in cui stanno in ogni offerta.
# I minorenni (franchigie 0-600, tetto di partecipazione 350, sconti famiglia)
# sono fuori dalla prima versione: vedi LEGGIMI.md.
FRANCHIGIE = [300, 500, 1000, 1500, 2000, 2500]
CHIAVI_FRANCHIGIA = {f"FRA-{f}": i for i, f in enumerate(FRANCHIGIE)}

# AKL-ERW = adulti dai 26 anni, AKL-JUG = giovani adulti 19-25.
# Stesse franchigie, stesso tetto di partecipazione (700 CHF): il calcolo e'
# identico, cambia solo il prezzo. AKL-KIN resta fuori.
CLASSI_ETA = {"AKL-ERW": "E", "AKL-JUG": "J"}

# I quattro tipi di modello, come li chiama l'UFSP.
TIPI_MODELLO = ["BASE", "HAM", "HMO", "DIV"]

# Nomi commerciali. Il registro federale contiene la ragione sociale completa
# ("Genossenschaft KRANKENKASSE SLKK"), che in una tabella di confronto e'
# rumore: qui la si riduce togliendo forme giuridiche e parole di mestiere.
# La riduzione automatica basta quasi sempre; queste sono le eccezioni in cui
# produce qualcosa di brutto o ambiguo.
NOMI_BREVI = {
    8: "CSS",
    290: "CONCORDIA",
    455: "ÖKK",
    509: "Vivao Sympany",
    780: "Glarner",
    923: "SLKK",
    1113: "Caisse-maladie d\u2019Entremont",
    1318: "Wädenswil",
    1479: "Mutuel",
    1507: "AMB",
    1535: "Philos",
    1542: "Assura",
}

# Parole che non distinguono un assicuratore dall'altro.
FORME_GIURIDICHE = [
    "société coopérative", "Genossenschaft", "Stiftung", "Verein", "AG", "SA",
]
PAROLE_MESTIERE = [
    "Kranken- und Unfallversicherungen", "Grundversicherungen",
    "Gesundheitsversicherung", "Krankenversicherungen", "Krankenversicherung",
    "Gesundheitsgruppe", "gesundheitsgruppe", "Versicherungen",
    "Krankenkasse", "KRANKENKASSE", "Assurance Maladie",
]


def nome_breve(numero: int, ragione: str) -> str:
    """Riduce la ragione sociale a come la gente chiama l'assicuratore."""
    if numero in NOMI_BREVI:
        return NOMI_BREVI[numero]
    s = ragione
    for parola in FORME_GIURIDICHE + PAROLE_MESTIERE:
        s = re.sub(rf"(?<!\w){re.escape(parola)}(?!\w)", " ", s)
    s = re.sub(r"\s+", " ", s).strip(" ,.-")
    return s or ragione


# --------------------------------------------------------------- utilita'

def pulisci(testo: object) -> str:
    """Normalizza una cella Excel: gli a capo dentro le celle del registro
    federale spezzano le parole, a volte con un trattino di sillabazione
    ("Krankenver-\\nsicherung") e a volte senza. Il trattino si toglie se la
    riga dopo comincia in minuscolo, si tiene se comincia in maiuscolo: nel
    secondo caso fa parte del nome ("Kranken-Versicherung")."""
    if testo is None:
        return ""
    s = str(testo).replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"-[ \t]*\n\s*([a-zäöüàéèç])", r"\1", s)
    s = re.sub(r"-[ \t]*\n\s*([A-ZÄÖÜ])", r"-\1", s)
    s = s.replace("\n", " ")
    s = unicodedata.normalize("NFC", s)
    return re.sub(r"\s+", " ", s).strip()


def num(v: object) -> float | None:
    if v in (None, ""):
        return None
    try:
        return round(float(str(v).replace(",", ".")), 2)
    except ValueError:
        return None


def scrivi(percorso: Path, dati: object) -> int:
    percorso.parent.mkdir(parents=True, exist_ok=True)
    testo = json.dumps(dati, ensure_ascii=False, separators=(",", ":"))
    percorso.write_text(testo, encoding="utf-8")
    return len(testo.encode("utf-8"))


def solo(schema: str) -> Path:
    """Trova un file sorgente per pattern, con un errore leggibile se manca."""
    trovati = sorted(FONTE.glob(schema))
    if not trovati:
        sys.exit(
            f"Manca il file sorgente '{schema}' in {FONTE}.\n"
            "Le istruzioni per scaricarlo stanno in LEGGIMI.md."
        )
    return trovati[-1]


# ------------------------------------------------------------ assicuratori

def leggi_assicuratori() -> dict[int, dict]:
    """Registro federale degli assicuratori ammessi: numero UFSP -> nome.

    Il foglio ha righe di continuazione senza numero, che proseguono il nome
    dell'assicuratore precedente. Si accumulano finche' non ricompare un
    numero."""
    percorso = solo("zugelassene-krankenversicherer-*.xlsx")
    ws = openpyxl.load_workbook(percorso, data_only=True, read_only=True)[
        "Zugelassene Krankenversicherer"
    ]

    assicuratori: dict[int, dict] = {}
    corrente: int | None = None
    pezzi: list[str] = []
    gruppi: dict[int, str] = {}

    def chiudi() -> None:
        if corrente is not None:
            # I pezzi si uniscono con un a capo, non con uno spazio: e' cosi'
            # che pulisci() riconosce la sillabazione di fine riga
            # ("Krankenver-" + "sicherung AG" -> "Krankenversicherung AG").
            assicuratori[corrente] = {
                "ragione_sociale": pulisci("\n".join(pezzi)),
                "gruppo": gruppi.get(corrente, ""),
            }

    for riga in ws.iter_rows(min_row=2, values_only=True):
        numero, _uid, nome, *resto = (list(riga) + [None] * 7)[:7]
        # Alcune righe hanno il numero seguito da un contrassegno ("1179 x"):
        # sono assicuratori che non offrono piu' l'assicurazione di base. Vanno
        # comunque riconosciuti come inizio di una nuova voce, altrimenti il
        # loro nome finisce appiccicato a quello dell'assicuratore precedente.
        testo_numero = str(numero).strip() if numero is not None else ""
        avvio = re.match(r"^(\d+)", testo_numero)
        if avvio:
            chiudi()
            corrente = int(avvio.group(1))
            pezzi = [str(nome)] if nome else []
            gruppo = pulisci(resto[2] if len(resto) > 2 else "")
            gruppi[corrente] = "" if gruppo in ("---", "") else gruppo
        elif corrente is not None and nome:
            pezzi.append(str(nome))
    chiudi()
    return assicuratori


# ------------------------------------------------------------------ modelli

def leggi_modelli() -> dict[str, dict]:
    """Tarife.csv: nome commerciale del modello in DE/FR/IT, per assicuratore.

    Il file e' separato da punto e virgola (il CSV dei premi da virgole) e il
    numero dell'assicuratore e' riempito di zeri: '0008' contro '8'."""
    percorso = FONTE / "Tarife.csv"
    if not percorso.exists():
        return {}
    modelli: dict[str, dict] = {}
    with percorso.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            if r.get("Kategorie") != "MOD":
                continue
            chiave = f"{int(r['Versicherer'])}|{r['Tarif']}"
            modelli[chiave] = {
                "tipo": r["Tariftyp"],
                "de": pulisci(r["Name_DE"]),
                "fr": pulisci(r["Name_FR"]),
                "it": pulisci(r["Name_IT"]),
            }
    return modelli


# --------------------------------------------------------------------- NAP

def leggi_nap() -> tuple[dict, dict]:
    """Foglio B_NPA di praemienregionen.xlsx: NAP -> cantone e regione.

    E' il pezzo che il CSV dei premi non ha. I premi sono indicizzati per
    cantone e regione di premio (PR-REG CH0..CH3), non per NAP: senza questa
    tabella l'utente dovrebbe sapere in che regione abita, che e' esattamente
    la cosa che non sa e non deve sapere.

    Un NAP puo' stare in piu' regioni (i comuni sparsi su piu' valli): in quel
    caso restituiamo tutte le possibilita' e il sito chiede quale sia la
    localita' giusta."""
    ws = openpyxl.load_workbook(solo("praemienregionen.xlsx"),
                                data_only=True, read_only=True)["B_NPA"]

    nap: dict[str, list] = defaultdict(list)
    regioni_per_cantone: dict[str, set] = defaultdict(set)
    visti: set[tuple] = set()
    intestazione_passata = False

    for riga in ws.iter_rows(min_row=1, values_only=True):
        celle = list(riga) + [None] * 8
        plz, ort, kanton, region = celle[1], celle[2], celle[3], celle[4]
        gemeinde = celle[6]
        if not intestazione_passata:
            # l'intestazione vera e' la riga che contiene 'PLZ'
            if plz and "PLZ" in str(plz):
                intestazione_passata = True
            continue
        if plz is None or kanton is None or region is None:
            continue
        codice = str(int(plz)) if isinstance(plz, (int, float)) else str(plz).strip()
        if not codice.isdigit():
            continue
        cantone = str(kanton).strip()
        regione = str(int(region)) if isinstance(region, (int, float)) else str(region).strip()
        localita = pulisci(ort) or pulisci(gemeinde)
        comune = pulisci(gemeinde)
        regioni_per_cantone[cantone].add(regione)
        chiave = (codice, localita, cantone, regione)
        if chiave in visti:
            continue
        visti.add(chiave)
        nap[codice].append([localita, cantone, regione, comune])

    return dict(nap), {k: sorted(v, key=int) for k, v in regioni_per_cantone.items()}


# ------------------------------------------------------------------- premi

def leggi_premi() -> tuple[dict, dict, int]:
    """Prämien_CH.csv -> offerte raggruppate per cantone e regione.

    Un'"offerta" e' una combinazione assicuratore + modello + copertura
    infortuni + classe d'eta', con la lista dei premi per le sei franchigie.
    Le franchigie che quell'assicuratore non offre restano None, e non e' un
    caso limite: su circa 9 900 combinazioni adulte, 500 hanno cinque
    franchigie invece di sei. Il vecchio Insurek pretendeva sempre sei numeri;
    qui semplicemente si valuta quel che c'e'."""
    percorso = FONTE / "Prämien_CH.csv"
    gruppi: dict[tuple, dict] = {}
    anno = 0

    with percorso.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            classe = CLASSI_ETA.get(r["Altersklasse"])
            if classe is None:
                continue
            indice = CHIAVI_FRANCHIGIA.get(r["Franchise"])
            if indice is None:
                continue
            anno = anno or int(r["Geschäftsjahr"])
            regione = r["Region"].replace("PR-REG CH", "")
            chiave = (r["Kanton"], regione,
                      int(r["Versicherer"]), r["Tarif"],
                      1 if r["Unfalleinschluss"] == "MIT-UNF" else 0,
                      classe)
            g = gruppi.get(chiave)
            if g is None:
                g = gruppi[chiave] = {
                    "premi": [None] * len(FRANCHIGIE),
                    "tipo": r["Tariftyp"].replace("TAR-", ""),
                    "base": r["isBaseP"] == "1",
                    "etichetta": pulisci(r["Tarifbezeichnung"]),
                }
            g["premi"][indice] = num(r["Prämie"])

    per_regione: dict[tuple, list] = defaultdict(list)
    etichette: dict[str, dict] = {}
    for (cantone, regione, ass, tariffa, unf, classe), g in gruppi.items():
        per_regione[(cantone, regione)].append(
            [ass, tariffa, classe, unf, g["premi"]]
        )
        chiave = f"{ass}|{tariffa}"
        if chiave not in etichette:
            etichette[chiave] = {"tipo": g["tipo"], "base": g["base"],
                                 "de": g["etichetta"]}
    return dict(per_regione), etichette, anno


# -------------------------------------------------------------------- main

def main() -> None:
    print("Insurek — costruzione dei dati\n")

    assicuratori = leggi_assicuratori()
    print(f"  assicuratori nel registro federale : {len(assicuratori)}")

    modelli_tarife = leggi_modelli()
    print(f"  modelli con nome tradotto          : {len(modelli_tarife)}")

    nap, regioni_per_cantone = leggi_nap()
    print(f"  NAP                                : {len(nap)}")

    per_regione, etichette, anno = leggi_premi()
    print(f"  anno di premio                     : {anno}")
    print(f"  aree (cantone + regione)           : {len(per_regione)}")

    # I modelli: nome tradotto da Tarife.csv dove c'e', altrimenti la
    # denominazione tedesca che sta nel CSV dei premi. Molti sono nomi
    # commerciali ("Callmed", "KPTwin.doc") e restano uguali in ogni lingua:
    # a distinguerli e' il tipo, che il sito traduce da solo.
    modelli: dict[str, dict] = {}
    for chiave, base in etichette.items():
        tradotto = modelli_tarife.get(chiave, {})
        nome_de = tradotto.get("de") or base["de"]
        modelli[chiave] = {
            "t": base["tipo"],
            "b": 1 if base["base"] else 0,
            "de": nome_de,
            "fr": tradotto.get("fr") or nome_de,
            "it": tradotto.get("it") or nome_de,
        }

    # Anagrafica assicuratori, limitata a quelli che compaiono nei premi.
    presenti = sorted({o[0] for offerte in per_regione.values() for o in offerte})
    anagrafica: dict[str, dict] = {}
    senza_nome_breve = []
    for numero in presenti:
        registro = assicuratori.get(numero, {})
        ragione = registro.get("ragione_sociale", f"Assicuratore {numero}")
        breve = nome_breve(numero, ragione)
        if numero not in assicuratori:
            senza_nome_breve.append((numero, ragione))
        anagrafica[str(numero)] = {"n": breve, "r": ragione,
                                   "g": registro.get("gruppo", "")}

    # Nel CSV dei premi compaiono un paio di codici di territorio ("ZE",
    # "ZR") che non sono cantoni e non stanno nella tabella delle regioni di
    # premio: nessun NAP porta li', quindi sarebbero aree irraggiungibili.
    fantasma = sorted({c for c, _ in per_regione} - set(regioni_per_cantone))
    if fantasma:
        for chiave in [k for k in per_regione if k[0] in fantasma]:
            del per_regione[chiave]
        print(f"  territori scartati (non sono cantoni) : {', '.join(fantasma)}")

    cantoni = sorted({c for c, _ in per_regione})
    meta = {
        "anno": anno,
        "generato": dt.date.today().isoformat(),
        "fonte": "Ufficio federale della sanità pubblica (UFSP) — opendata.swiss",
        "franchigie": FRANCHIGIE,
        "tetto_partecipazione": 700,
        "quota_partecipazione": 0.10,
        "cantoni": {c: regioni_per_cantone[c] for c in cantoni},
        "assicuratori": anagrafica,
        "modelli": modelli,
        "tipi": TIPI_MODELLO,
    }

    totale = 0
    totale += scrivi(USCITA / "meta.json", meta)
    totale += scrivi(USCITA / "nap.json", nap)

    # Le aree cambiano da un anno all'altro (fusioni di comuni, assicuratori
    # che si ritirano da un cantone). I file dell'anno prima che nessuno
    # riscrive vanno tolti, altrimenti restano li' a mentire.
    prima = {f.name for f in (USCITA / "premi").glob("*.json")}
    scritti: set[str] = set()

    massimo = 0
    for (cantone, regione), offerte in sorted(per_regione.items()):
        offerte.sort(key=lambda o: (o[0], o[1], o[2], o[3]))
        nome = f"{cantone}-{regione}.json"
        scritti.add(nome)
        peso = scrivi(USCITA / "premi" / nome,
                      {"k": cantone, "r": regione, "a": anno, "o": offerte})
        massimo = max(massimo, peso)
        totale += peso

    residui = sorted(prima - scritti)
    for nome in residui:
        try:
            (USCITA / "premi" / nome).unlink()
        except OSError:
            print(f"  ATTENZIONE: non riesco a togliere sito/dati/premi/{nome}, "
                  "e' un'area che quest'anno non esiste piu': va tolta a mano.")

    print(f"\n  offerte totali                     : "
          f"{sum(len(v) for v in per_regione.values())}")
    print(f"  file dell'area piu' grande         : {massimo / 1024:.0f} KB")
    print(f"  peso complessivo di sito/dati/     : {totale / 1024:.0f} KB")
    print(f"  (il browser ne scarica meta.json + nap.json + un'area sola)")

    if senza_nome_breve:
        print("\n  Assicuratori presenti nei premi ma assenti dal registro "
              "(controllare la versione del file):")
        for numero, ragione in senza_nome_breve:
            print(f"    {numero}: {ragione}")

    print("\nFatto.")


if __name__ == "__main__":
    main()
